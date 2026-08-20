import docx
from docx.document import Document as _Doc
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
import sys, io, os

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

path = "../ANALIZA PRIJENOSA BOJE KOD VIŠEBOJNIH REPRODUKCIJA S UV-SUŠEĆIM I KONVENCIONALNIM TISKARSKIM BOJAMA_Sara_Iva_Merlić_v2.docx"
doc = docx.Document(path)

rels = doc.part.rels
rid_to_image = {}
for rid, rel in rels.items():
    if "image" in rel.reltype:
        rid_to_image[rid] = os.path.basename(rel.target_ref)

def para_images(p):
    imgs = []
    for blip in p._element.iter(qn('a:blip')):
        embed = blip.get(qn('r:embed'))
        if embed and embed in rid_to_image:
            imgs.append(rid_to_image[embed])
    return imgs

def iter_block_items(parent):
    if isinstance(parent, _Doc):
        parent_elm = parent.element.body
    else:
        parent_elm = parent._element
    for child in parent_elm.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, parent)
        elif child.tag == qn('w:tbl'):
            yield Table(child, parent)

idx = 0
for block in iter_block_items(doc):
    if isinstance(block, Paragraph):
        style = block.style.name if block.style else ''
        text = block.text.strip()
        imgs = para_images(block)
        if not text and not imgs:
            idx += 1
            continue
        print(f"[P{idx:04d}] [{style}] {text}")
        for im in imgs:
            print(f"    <IMG> {im}")
        idx += 1
    elif isinstance(block, Table):
        print(f"[T{idx:04d}] <TABLE rows={len(block.rows)} cols={len(block.columns)}>")
        for ri, row in enumerate(block.rows):
            cells = []
            for cell in row.cells:
                txt = " / ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
                # collect images inside cells
                cimgs = []
                for p in cell.paragraphs:
                    cimgs.extend(para_images(p))
                if cimgs:
                    txt += " [IMG:" + ",".join(cimgs) + "]"
                cells.append(txt)
            print(f"  R{ri}: | " + " | ".join(cells) + " |")
        print(f"[/TABLE]")
        idx += 1
